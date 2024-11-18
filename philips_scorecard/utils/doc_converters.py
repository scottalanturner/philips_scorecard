import base64
from io import BytesIO
from docx import Document
import io
import pandas as pd

def word_to_base64(file_path : str) -> str:
    """
    Converts a Word document to base64 string for HTTP transmission
    
    Parameters:
        file_path (str): Path to the Word document
        
    Returns:
        str: Base64 encoded string of the document
    """
    try:
        # Create a bytes buffer
        buffer = io.BytesIO()
        
        # Load and save the document to the buffer
        doc = Document(file_path)
        doc.save(buffer)
        
        # Get the bytes value and encode to base64
        doc_bytes = buffer.getvalue()
        base64_encoded = base64.b64encode(doc_bytes).decode('utf-8')
        
        return base64_encoded
        
    except Exception as e:
        raise Exception(f"Error processing document: {str(e)}")
    finally:
        buffer.close()

def excel_to_base64(file_path: str) -> str:
    """
    Converts an Excel file to base64 string for HTTP transmission
    
    Parameters:
        file_path (str): Path to the Excel file
        
    Returns:
        str: Base64 encoded string of the Excel file
    """
    try:
        # Create a bytes buffer
        buffer = io.BytesIO()
        
        # Load and save the Excel file to the buffer
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            # Read all sheets from the Excel file
            excel_file = pd.ExcelFile(file_path)
            # Copy each sheet to the buffer
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        # Get the bytes value and encode to base64
        excel_bytes = buffer.getvalue()
        base64_encoded = base64.b64encode(excel_bytes).decode('utf-8')
        
        return base64_encoded
        
    except Exception as e:
        raise Exception(f"Error processing Excel file: {str(e)}")
    finally:
        buffer.close()

def convert_doc_to_base64(document : Document) -> str:
    # Save updated document to a BytesIO buffer
    output = BytesIO()
    document.save(output)
    output.seek(0)

    # Encode modified document to base64. This would return in the HTTP request normally
    content = base64.b64encode(output.read()).decode("utf-8")

    return content


def get_document(document_content_base64):
    # The base64 content of the Word document is transmitted in the HTTP Post
    # It then has to be decoded, and then the placeholders can be replaced
    document_content = base64.b64decode(document_content_base64)
    document = Document(BytesIO(document_content))
    return document

def convert_base64_to_excel_sheets(base64_content: str) -> dict:
    """
    Takes a base64 encoded Excel file and reads it with pandas
    
    Parameters:
    base64_content (str): The base64 string from your JSON
    
    Returns:
    dict: Dictionary of all sheets in the Excel file (dictionary of DataFrames)
    """
    try:
        # Decode base64 to bytes
        excel_bytes = base64.b64decode(base64_content)
        
        # Create a BytesIO object (in-memory file)
        excel_buffer = io.BytesIO(excel_bytes)
        
        # Read Excel file using pandas
        sheets = pd.read_excel(excel_buffer, sheet_name=None)
        
        return sheets
    except Exception as e:
        raise Exception(f"Error reading Excel from base64: {str(e)}")
